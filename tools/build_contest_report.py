"""Fill retained official template; no submission or credential inclusion."""
from pathlib import Path
from copy import deepcopy
import hashlib, json
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

OUT=Path(__file__).resolve().parents[1]/'deliveries/contest-report-20260917'
REF=OUT/'官方作品提交模板.docx'
d=Document(REF)
paras=list(d.paragraphs)
def fill(p,text):
    props=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear(); r=p.add_run(text)
    if props is not None:r._r.insert(0,props)
    return p

content={
0:'2026 首届 openvela AI 硬件开发者大赛\nVelaVision 技术报告',
7:'VelaVision 面向家庭与门店的标准化皮肤影像采集，基于瑞芯微 RK3576 KICKPI K7 开展 openvela 原生适配，集成蓝牙配网、流式语音命令、人脸跟随、三角度拍摄和后端报告播报。K7 负责视觉、姿态判定及多媒体采集播放，STM32 执行云台动作，开发阶段网关连接云端语音与业务服务。已验证三张实拍图上传受理；一次30分钟视频会话接收53,805帧、29.88 FPS。固定提示语存于板端，减少对云TTS的依赖。本作品当前为可联调原型，尚依赖主机网关，临床有效性与长期稳定性未验收。',
10:'项目背景与问题定义：同一人的正面、左侧和右侧照片若拍摄角度不一致，会影响后续观察与对比。手持拍摄需要反复调整角度，且使用者难以同时看屏幕、转头和操作按钮。我们以“语音发起、云台跟随、角度达标后自动采集”为交互方式，降低多角度影像采集的操作负担。产品定位为皮肤影像采集与报告展示原型，不将未经验证的分数或建议作为医疗诊断。',
11:'技术难点：第一，RK3576 K7 需要打通原生启动、UART、USB 摄像头、音频及无线共享链路；第二，视频持续传输与流式录音、播放、蓝牙及Wi-Fi并发时，需要控制缓冲占用并恢复硬件状态；第三，语音识别、云台动作、姿态稳定判定与阶段提示具有不同的时序，必须保证取消有效、照片与判定帧一致、提示不会触发错误动作。',
12:'创新点：将硬件平台适配与应用闭环一起验证；以板端确定性命令状态机控制真实外设，而非由云端文字直接驱动电机；将固定提示语固化为本地PCM，把云TTS留给动态报告；以正面、左侧、右侧的目标锁定和质量门控完成三视角采集，并用任务ID追踪上传、分析和播报。创新点属于工程与交互集成，不主张创造新的医学诊断算法。',
14:'系统总体架构：手机App通过BLE向K7请求Wi-Fi列表并提交连接信息。K7运行openvela固件，负责麦克风采样、摄像头采集、原生人脸/姿态计算和云台控制；STM32运行已有云台固件，接收UART目标坐标并控制舵机。Windows语音网关负责云服务鉴权与协议转换，业务后端管理照片、检测任务及报告。Ubuntu窗口接收预览和日志，作为开发调试端，不承担板端跟踪判定。图1展示节点与数据流。',
15:'方案论证与选型：选择K7是为了验证ARM64平台上的多媒体与无线适配，并保留CPU、DDR及NPU的后续扩展空间；当前应用不把板卡标称八核、4GB DDR或6TOPS等同为已充分利用的性能。相较于把视觉完全放在PC，板端计算减少对预览链路的依赖；相较于本地运行所有语音模型，现阶段云端流式ASR更便于快速验证交互，代价是依赖网络及网关。断网时本地固定提示仍可用，云ASR/动态报告不能继续；失败不会伪造成功结果。完整离线语音、自动断线重试和独立联网TLS链路属于后续工作。',
16:'关键模块设计：①BLE配网服务处理App发送的X扫描请求、Wi-Fi列表、凭据提交与联网状态事件；②音频使用16kHz、单声道、S16LE PCM，按20ms帧发送流式ASR；③视觉流水线隔离USB采集、图像解码、检测/姿态和预览；④命令执行只接受受控意图，运行前确认模式切换结果；⑤上传网关使用设备Bearer会话处理三图及报告请求，凭据留在私有配置，文档与公开日志不包含密钥。',
18:'AI算法实现：板端人脸检测采用项目中的YuNet原生实现，姿态模块输出偏航、俯仰、横滚角，用于拍照条件判定。对应实现位于app/k7host的检测与姿态模块，使用CPU推理；当前报告不把独立NPU INT8矩阵诊断结果写成整机模型已获NPU加速。云端流式ASR接入阿里云百炼paraformer-realtime-v2；speech-service现有TTS配置为qwen-audio-3.0-tts-flash。MiMo曾作为接口探索方向，本次验收链路不声称使用MiMo。模型API通过服务端配置鉴权，云台业务接口采用设备Bearer token。',
19:'关键机制设计：人脸框经目标确认和误差滤波后计算横纵坐标修正，设置死区、积分限幅及最大步长，避免单帧噪声引发大幅跳动。姿态采集先锁定单人目标，再检查角度、稳定性和图像质量：正脸偏航±7°，左右侧脸以±45°为目标、容差±13°；俯仰/横滚门限12°、偏航抖动门限3°、稳定保持500ms。镜像仅用于界面显示，不更改算法左右定义。三图完成后退出采集状态并保持云台停止。',
20:'openvela系统能力的运用：以openvela SDK及其应用构建体系为基础，新增K7 BSP、应用组件和驱动适配；通过USB Host/UVC、SAI/codec、原生音频采集播放、Bluetooth及网络栈实现多媒体与通信，通过板端CPU模型实现AI感知。当前明确落地的是AI与多媒体，GTK预览窗口运行在Ubuntu，不计作板端图形能力。源码将新增应用、板级代码和SDK补丁分别置于app、board、port；提交前还需逐项列明所依赖的openvela仓库与组件，便于评审核查“基于openvela”的规则。没有把NuttX内核适配本身作为满足该规则的全部依据。',
22:'软件/固件架构：统一仓库VelaVision中，app/k7host管理视觉与姿态采集，app/k7agent/cloud管理云语音命令和报告协议，app/k7sound管理音频硬件与PCM，app/k7radio管理无线；board/kickpi_k7和port保存板级及SDK移植代码；host提供调试窗口和网关，mcu保留云台固件，frontend保存配网协议及联调说明。Ubuntu SDK副本通过tools/sync_sdk.py审计后同步，未知改动不被静默覆盖。固件采用Fastboot下载到RAM并校验CRC；当前版本未固化到eMMC。',
23:'数据流与关键流程：配网成功后，短音提示进入录音；ASR结果匹配“你好，openvela”后请求跟踪模式，模式应用成功再播放本地启动提示。“皮肤检测”进入三视角采集，依次播报对准及拍摄完成提示。三图齐备后停止跟踪、提交照片，收到taskId后轮询任务；queued/analyzing继续等待，needs_retake按requiredViews提示补拍，failed返回失败状态，report_ready使用reportId读取brief报告。409 TASK_REPLACED终止旧任务播报。补拍提示路径已实现，完整补拍重新提交流程尚需进一步验收。',
24:'硬件设计与适配：已开展并完成多项RK3576 KICKPI K7原生驱动适配。主板配置为4GB DDR、32GB级eMMC，外接USB UVC摄像头、STM32双轴云台及板载麦克风/喇叭。UART6以115200 8N1连接STM32，板端设备名/dev/ttyS1；调试UART0为1500000波特率。云台接线为K7 Pin5 TX→MCU RX、Pin7 RX←MCU TX、Pin6 GND。USB Host接相机，OTG用于RAM加载/预览传输。BOM数量见表2；摄像头和STM32具体采购型号、单价需队伍按实物补齐，不以推测填报。',
25:'应用与交互端设计：手机App负责BLE发现、通知订阅、Wi-Fi选择和凭据提交，成功后显示联网状态；凭据不经过语音输入，减少口述密码误识别。Ubuntu整机窗口显示镜像视频、人脸框、姿态角、三图状态和上传状态，提供相机待命、单次语音检查、整机语音测试及停止入口。按钮由新鲜板端状态驱动。界面与网关属于开发工具，当前未完成脱离电脑的独立交付。',
26:'自定义Skill：沉淀skills/k7-openvela-driver-port/SKILL.md，覆盖原厂资料对照、分层驱动移植、RAM验证、源码同步和证据归档。其作用是约束AI开发过程，要求把寄存器读回、协议成功、实际动作和人工验收分开记录，并保留失败与恢复证据。它是开发期Skill；当前没有部署/data/agent/skills/运行时Agent Skill，不将本地规则命令解释器冒充openvelaClaw。若按AI硬件产品创新方向申报运行时Skill能力，该项仍需补做；本报告主要以新硬件适配和应用验证呈现。',
28:'测试环境：KICKPI K7真机、USB摄像头、STM32云台、板载麦克风/喇叭；Windows主机运行speech-service与语音网关，VMware Ubuntu提供SDK构建和串口/USB联调。量化结果取自仓内串口日志、网关受理记录及构建校验，未使用模拟图像替代三图实拍。不同测试版本分开列示，不将旧版无线长测直接视为新版整机稳定性结论。',
29:'功能测试：手机BLE扫描、获取Wi-Fi列表、提交凭据及联网状态回传已通过真人联调；“你好，openvela”启动跟踪，用户于2026-09-17明确确认跟随已修复；三角度拍摄及真实照片上传已获得后端受理。固定阶段提示在板端播放。报告SSE分段播放此前使用mock文案验证，用户确认完整听到；现已改为真实任务brief内容，后端目前缺少详细说明字段，新模式板端完整实听尚待确认。以上分别记录“用户确认”“协议成功”和“待验收”，不互相替代。',
30:'性能测试：表3列出可复核实测；预览接收帧率、板端解码帧率和模型推理帧率不是同一个指标。语音优化的400ms是静音收尾阈值，不是“说完到电机动作/喇叭出声”的完整延迟。当前没有足够样本给出端到端P50/P95，亦未测量功耗、续航或正式识别准确率。',
31:'可靠性与稳定性：历史30分钟Wi-Fi测试1500发1499收，零丢包标准未通过；另一版BLE/Wi-Fi共存记录1497/1500且BLE保持通过，仍不能视为生产级稳定。当前视频会话1800秒后按配置退出，同一次启动不能直接重复初始化，需重启恢复。异常路径包括USB透传断开、串口I/O错误、BLE通知订阅失败、TCP发送缓冲不足，均保留诊断与恢复记录。云台协议没有MCU位置回执，UART发送成功不能证明运动完成，实际跟随以用户观察为证。未开展人群样本准确率、误报/漏报统计或皮肤分析临床验证；目标丢失、角度/质量不达标时不进入正常完成流程，网络失败不伪造报告。',
33:'AI开发过程与效率：我们使用Codex开展驱动阅读、补丁生成、单元测试、构建同步、串口日志分析及报告整理。典型问题包括“日志发送成功但电机未确认”“音频声道配置导致回放无声”“大图上传受缓冲限制”；处理方式是保留失败日志，拆分驱动、协议、云端和人工验收层级，再用固定版本回归。AI提升的是检索、实现和验证迭代效率，未做有/无AI的人时对照，因此不提供虚构的倍数。19个导出日志文件、43,295条事件已通过原版官方校验器；该数是事件数，不是Token总量。真实会话范围、脱敏与采集不足保留在同仓说明。',
35:'成果总结：完成K7上openvela原生启动与多项外设链路接入，在同一应用中串联蓝牙配网、云流式语音命令、板端视觉跟随、三角度采集和后端上传；固定提示语由板端播放，动态报告由网关按后端内容分段合成。用户确认实际跟随已恢复。项目以可复核代码、固件哈希、真实日志和操作反馈作为成果依据，已构建、已部署与已验收状态分开维护。',
36:'应用前景与商业价值：目标用户为需要标准化多角度皮肤影像记录的成年个人、护肤门店及研发演示人员；优先验证室内固定机位、中文语音与手机配网场景。拟采用设备与后端服务组合的交付方式，价值来自降低重复拍摄操作和建立一致的影像记录。尚无销售、用户规模、收入或临床结论数据；商业化需进一步解决隐私授权、部署稳定性、成本与后端分析有效性。',
37:'不足与未来工作：①将主机网关迁移到稳定服务或完成板端独立协议栈，明确断网降级；②实现可重复开启的长时相机会话和持久化启动；③完成运行时Skill及合规指定唤醒词的严格匹配，当前“你好”前缀兼容用于调试，正式参赛演示使用“你好，openvela”；④量化连续识别准确率、端到端延迟、功耗和跨设备并发；⑤将后端真实报告说明和正式流式讲解接通并验收；⑥整理第三方模型/代码许可、复现步骤及代码日志对应关系。新模型、八核并发、完整NPU推理和医疗效果不列为当前已完成成果。'
}
for i,t in content.items():fill(paras[i],t)

# Fill only official editable cells; retain all original tables and guidance.
for i,t in {1:'VelaVision 基于 openvela 的语音交互三视角皮肤影像采集系统',2:'待队长确认正式队伍名称（仓库标识：yuanxinshixisheng）',3:'待队长补充各成员姓名与实际分工；不得将AI工具列为团队成员。',4:'新硬件平台适配＋AI硬件产品创新应用验证（最终报名方向由队长核对）'}.items():fill(d.tables[2].cell(i,1).paragraphs[0],t)
stats={1:'未统计可靠比例；源码与AI修改轨迹可核查，不以日志事件数估算代码占比。',2:'Codex（开发、测试、日志与文档）；云端语音模型用于产品运行，另行说明。',3:'使用文件/终端、浏览器及文档工具核查资料与执行验证；未声称使用VelaJS MCP或Figma MCP。',4:'新增k7-openvela-driver-port开发期Skill，路径skills/k7-openvela-driver-port/SKILL.md；未部署运行时Agent Skill。',5:'未完成统一统计，待汇总控制台账单；43,295为日志事件数，不是Token数。'}
for i,t in stats.items():fill(d.tables[3].cell(i,1).paragraphs[0],t)

def insert_par(after,text,bullet=False):
    e=deepcopy(paras[10]._p); after._p.addnext(e)
    from docx.text.paragraph import Paragraph
    p=Paragraph(e,after._parent);fill(p,text)
    if not bullet:
        pr=p._p.pPr
        n=pr.find(qn('w:numPr'))
        if n is not None:pr.remove(n)
    return p

info=insert_par(paras[5],'编制日期：2026年9月17日。状态：技术内容审阅稿，团队信息、最终视频与远端提交状态待补齐。')
insert_par(info,'源码工作仓：https://github.com/allenxun/contest2026_483_yuanxinshixisheng；目标分支：dev-ai-contest-2026。当前产品代码尚有未提交内容，最终提交前须完成专属仓合入及远端核验。')

def table_after(p,headers,rows,widths=None):
    t=d.add_table(rows=1,cols=len(headers))
    src=d.tables[2]._tbl.tblPr
    t._tbl.remove(t._tbl.tblPr);t._tbl.insert(0,deepcopy(src))
    p._p.addnext(t._tbl)
    widths=widths or [1.12,2.22,2.26]
    t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for c,text in zip(t.rows[0].cells,headers):c.text=text
    for row in rows:
        for c,text in zip(t.add_row().cells,row):c.text=str(text)
    for j,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if j==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci])
            for p0 in c.paragraphs:
                p0.paragraph_format.space_after=Pt(4);p0.paragraph_format.space_before=Pt(4)
                for r in p0.runs:
                    r.font.size=Pt(10);r.font.name='Arial';r.bold=j==0;r._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'等线')
    return t

table_after(insert_par(paras[24],'表2　原型硬件组成'),['部件','数量与用途','接口/状态'],[
 ('KICKPI K7 / RK3576','1；openvela主处理板','USB、UART、无线与板载音频'),
 ('USB UVC摄像头','1；人脸及三视角影像','USB Host；实物型号待补'),
 ('STM32控制板与双轴云台','1套；舵机执行','UART6；已有v1.3-xcenter协议'),
 ('麦克风、codec与喇叭','板载音频链路','16kHz单声道PCM'),
 ('调试与供电配件','USB串口、OTG线、独立供电','用于RAM加载、日志与执行器供电')])
metrics=[
 ('视频接收','53,805帧 / 1800.226秒；29.88 FPS','command-latency版；UVC会话，不是模型FPS'),
 ('板端解码','11.86 FPS；平均3915μs/帧','同一会话；解码耗时不含全部推理与网络'),
 ('固件RAM加载','21,784,736字节 / 2.302秒','2026-09-17重启加载；固件CRC 654ca198'),
 ('三图实拍上传','正26344B／左25785B／右26124B','后端accepted=true、queued；2026-09-17'),
 ('语音提前收尾','单次录音1740ms；最后活跃1340ms','真实命令日志；差值400ms，非端到端延迟'),
 ('侧脸条件','目标±45°，容差±13°；稳定500ms','配置与边界测试；非实测角度精度'),
 ('历史Wi-Fi长测','1500发1499收，约30分钟','2026-09-09旧版；严格零丢包未通过'),
 ('报告相关回归','19项自动测试通过','真实brief切换；新模式喇叭实听待确认')]
table_after(insert_par(paras[30],'表3　量化结果与适用范围'),['测试项','实测/配置结果','版本与边界'],metrics)
insert_par(paras[31],'证据索引：E1 evidence/command-latency-20260917/camera-session-ended.log；E2 同目录restart-1352/startup.log与verification.json；E3 evidence/photo-upload-accepted-20260917/verification.json；E4 evidence/command-latency-20260917/tracking-investigation.log及tracking-user-acceptance.json；E5 docs/侧脸容差调整_20260917.md；E6 docs/无线稳定性相关记录与evidence/stability-20260909/；E7 docs/真实报告内容播报切换_20260917.md。原图与认证信息位于private，公开提交须使用经授权、脱敏的证据。')

# Diagrams are deterministic engineering drawings, not generated bitmap art.
fontpath='C:/Windows/Fonts/msyh.ttc'
def diagram(filename,boxes):
    im=Image.new('RGB',(1100,120+len(boxes)*122),'white');dr=ImageDraw.Draw(im)
    ft=ImageFont.truetype(fontpath,27)
    for i,(title,sub) in enumerate(boxes):
        y=30+i*122;dr.rounded_rectangle((35,y,1065,y+92),radius=10,fill='#f1f4f8',outline='#667788',width=2)
        dr.text((58,y+10),title.replace('↔','双向连接'),font=ft,fill='black');dr.text((58,y+48),sub.replace('↔','双向连接'),font=ft,fill='#334455')
        if i<len(boxes)-1:
            dr.line((550,y+93,550,y+119),fill='#445566',width=3);dr.polygon([(543,y+111),(557,y+111),(550,y+120)],fill='#445566')
    im.save(OUT/filename)
def figure(after,name,caption):
    p=insert_par(after,'');p.add_run().add_picture(str(OUT/name),width=Inches(5.65));p.paragraph_format.keep_with_next=True
    cp=insert_par(p,caption);cp.paragraph_format.space_after=Pt(8)
diagram('architecture.png',[
 ('手机 App → BLE → K7','发现设备、X请求Wi-Fi列表、提交凭据、接收联网事件'),
 ('K7 运行 openvela','USB摄像头、CPU视觉/姿态、录音/播放、本地提示与命令状态机'),
 ('执行与调试支路','UART6 ↔ STM32云台；USB预览 → Ubuntu联调窗口'),
 ('K7 Wi-Fi ↔ Windows语音网关','流式PCM上传、报告音频下发；云端密钥保留服务端'),
 ('云端语音服务 / 业务后端','ASR与TTS；三图受理 → taskId → reportId → brief报告')])
figure(paras[14],'architecture.png','图1　当前开发原型的节点与数据流。网关和预览仍依赖主机。')
diagram('flow.png',[
 ('App配网 → K7确认联网','保持云台停止，现场确认后启动相机待命'),
 ('短音 → 流式ASR → 受控意图','你好，openvela → 跟踪；皮肤检测 → 三角度采集'),
 ('正脸 → 左侧 → 右侧','角度/质量/稳定条件满足后保存对应帧并本地提示'),
 ('三图完成 → 云台停止 → 后端上传','取得taskId；未受理不宣称上传成功'),
 ('轮询任务 → report_ready → reportId','needs_retake提示补拍；failed或TASK_REPLACED停止旧流程'),
 ('读取brief → 逐句TTS → K7播放','每段播放确认；不将固定mock文案当真实报告')])
figure(paras[23],'flow.png','图2　语音控制、三图采集及报告播报流程。')

# Keep headings and table rows with their following content without changing page geometry.
for p in d.paragraphs:
    p.paragraph_format.widow_control=True
    p.paragraph_format.keep_together=True
    if p.text.startswith(('3.1 ','3.2 ','3.3 ','3.4 ','3.5 ','3.6 ','3.7 ','一、','二、','三、','四、')) or p.text in ['1、信息表','2、摘要','3、正文']:
        p.paragraph_format.keep_with_next=True
for t in d.tables:
    hp=t.rows[0]._tr.get_or_add_trPr()
    if hp.find(qn('w:tblHeader')) is None:hp.append(OxmlElement('w:tblHeader'))
    for p in t.rows[0].cells[0].paragraphs:p.paragraph_format.keep_with_next=True
    for row in t.rows:
        pr=row._tr.get_or_add_trPr()
        if pr.find(qn('w:cantSplit')) is None:pr.append(OxmlElement('w:cantSplit'))
dst=OUT/'VelaVision_官方模板技术报告_审阅稿.docx';d.save(dst)
# python-docx normalizes opaque XML on save; retain the original template parts byte-for-byte.
import zipfile
with zipfile.ZipFile(REF) as z:
    retained={n:z.read(n) for n in z.namelist() if any(k in n for k in ['header','footer','numbering','theme'])}
temp=dst.with_suffix('.tmp')
with zipfile.ZipFile(dst) as z,zipfile.ZipFile(temp,'w',zipfile.ZIP_DEFLATED) as w:
    for n in z.namelist():w.writestr(n,retained.get(n,z.read(n)))
temp.replace(dst)
(OUT/'待补信息.txt').write_text('正式提交前补齐：\n1. 队伍正式名称、成员姓名与分工。\n2. AI代码占比及Token用量若无法统计，保留未统计说明。\n3. 摄像头、STM32实物型号及必要的BOM信息。\n4. 专属仓最终地址、提交哈希与远端日志核验。\n5. 不超过5分钟的演示视频和硬件多角度照片。\n6. 新brief报告播报的实听确认。\n7. 核对openvela非NuttX仓库能力证据、运行时Skill要求、正式唤醒词和第三方许可。\n',encoding='utf-8')
print(dst)
